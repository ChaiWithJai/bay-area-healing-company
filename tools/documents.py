"""Local, provenance-preserving document IO. JSON request/response on stdio."""
import csv
import hashlib
import io
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
from xml.sax.saxutils import escape


def digest(data):
    return hashlib.sha256(data).hexdigest()


def stringify(value):
    return '' if value is None else str(value)


def read_sources(root):
    root = Path(root).resolve(strict=True)
    if not root.is_dir():
        raise ValueError('Input must be a directory')
    sources, errors = [], []
    for path in sorted(root.rglob('*')):
        if path.is_dir():
            if path.is_symlink():
                errors.append({'path': str(path.relative_to(root)), 'error': 'Symlink directory not traversed'})
            continue
        relative = path.relative_to(root).as_posix()
        try:
            resolved = path.resolve(strict=True)
            if not resolved.is_relative_to(root) or path.is_symlink():
                raise ValueError('Symlink input is not permitted')
            raw = path.read_bytes()
            kind = path.suffix.lower().lstrip('.')
            anchors, sections, tables = [], [], None
            if kind in ('txt', 'md', 'json', 'csv', 'tsv'):
                text = raw.decode('utf-8-sig')
                if kind == 'json':
                    json.loads(text)
                if kind in ('csv', 'tsv'):
                    rows = list(csv.reader(io.StringIO(text), delimiter='\t' if kind == 'tsv' else ','))
                    tables = {'sheets': [{'name': path.stem, 'headers': rows[0] if rows else [], 'rows': rows[1:]}]}
                anchors = [{'line': i + 1, 'text': line} for i, line in enumerate(text.splitlines())]
            elif kind == 'docx':
                from docx import Document
                from docx.table import Table
                from docx.text.paragraph import Paragraph
                document = Document(path)
                for i, element in enumerate(document.element.body):
                    if element.tag.endswith('}p'):
                        value = Paragraph(element, document).text
                    elif element.tag.endswith('}tbl'):
                        value = '\n'.join('\t'.join(cell.text for cell in row.cells) for row in Table(element, document).rows)
                    else:
                        continue
                    sections.append(value)
                    anchors.append({'block': i + 1, 'text': value})
                text = '\n'.join(sections)
            elif kind == 'xlsx':
                from openpyxl import load_workbook
                workbook = load_workbook(path, data_only=False, read_only=True)
                sheets = []
                for sheet in workbook:
                    rows = []
                    for row in sheet.iter_rows():
                        values = []
                        for cell in row:
                            value = cell.value
                            if hasattr(value, 'isoformat'):
                                value = value.isoformat()
                            values.append(value)
                            if value is not None:
                                anchors.append({'sheet': sheet.title, 'cell': cell.coordinate, 'text': stringify(value)})
                        rows.append(values)
                    sheets.append({'name': sheet.title, 'headers': rows[0] if rows else [], 'rows': rows[1:]})
                    sections.append('Sheet: ' + sheet.title + '\n' + '\n'.join('\t'.join(map(stringify, row)) for row in rows))
                workbook.close()
                text = '\n\n'.join(sections)
                tables = {'sheets': sheets}
            elif kind == 'pdf':
                from pypdf import PdfReader
                reader = PdfReader(path)
                for i, page in enumerate(reader.pages):
                    value = page.extract_text() or ''
                    method = 'text'
                    if not value.strip():
                        if not shutil.which('pdftoppm') or not shutil.which('tesseract'):
                            raise ValueError(f'Page {i+1} requires OCR: install local pdftoppm and tesseract')
                        with tempfile.TemporaryDirectory(prefix='wm-ocr-') as temporary:
                            prefix = str(Path(temporary) / 'page')
                            subprocess.run(['pdftoppm', '-f', str(i+1), '-l', str(i+1), '-singlefile', '-r', '200', '-png', str(path), prefix], check=True, capture_output=True, timeout=120)
                            value = subprocess.run(['tesseract', prefix + '.png', 'stdout'], check=True, capture_output=True, text=True, timeout=120).stdout
                        method = 'ocr'
                        if not value.strip():
                            raise ValueError(f'Page {i+1} has no readable content after OCR')
                    anchors.append({'page': i + 1, 'method': method, 'text': value})
                    sections.append(value)
                text = '\n\n'.join(sections)
            else:
                raise ValueError(f'Unsupported input format: {kind or "(none)"}')
            source = {'id': 'src-' + digest(relative.encode() + b'\0' + raw)[:20], 'path': relative, 'text': text, 'sha256': digest(raw), 'kind': kind, 'anchors': anchors}
            if tables is not None:
                source['tables'] = tables
            sources.append(source)
        except Exception as error:
            errors.append({'path': relative, 'error': str(error)})
    return {'sources': sources, 'errors': errors}


def safe_target(root, relative):
    relative = Path(relative)
    if relative.is_absolute() or '..' in relative.parts or str(relative) in ('', '.'):
        raise ValueError('Artifact path must remain within output directory')
    target = root / relative
    current = root
    for part in relative.parts:
        current = current / part
        if current.is_symlink():
            raise ValueError('Artifact path cannot contain symlinks')
    if not target.resolve().is_relative_to(root):
        raise ValueError('Artifact path escapes output directory')
    return target


def render_artifacts(directory, artifacts):
    requested = Path(directory).absolute()
    if any(parent.is_symlink() for parent in (requested, *requested.parents)):
        raise ValueError('Output directory cannot contain symlinks')
    requested.mkdir(parents=True, exist_ok=True)
    root = requested.resolve()
    targets = []
    for artifact in artifacts:
        target = safe_target(root, artifact['path'])
        if target in targets:
            raise ValueError('Duplicate artifact path')
        targets.append(target)
    rendered = []
    for artifact, target in zip(artifacts, targets):
        kind = artifact['kind']
        paragraphs = artifact.get('paragraphs', [])
        tables = artifact.get('tables', [])
        target.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(dir=target.parent, suffix='.' + kind, delete=False) as temporary:
            temporary_path = Path(temporary.name)
        try:
            if kind == 'docx':
                from docx import Document
                from docx.shared import RGBColor
                document = Document()
                for style_name in ('Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3'):
                    style = document.styles[style_name]
                    style.font.color.rgb = RGBColor(0, 0, 0)
                    style.font.underline = False
                    for border in style.element.xpath('./w:pPr/w:pBdr'):
                        border.getparent().remove(border)
                if artifact.get('title'):
                    document.add_heading(artifact['title'], 0)
                for paragraph in paragraphs or [artifact.get('text', '')]:
                    document.add_paragraph(str(paragraph))
                for table in tables:
                    headers = table['headers']
                    native = document.add_table(rows=1, cols=len(headers))
                    for cell, value in zip(native.rows[0].cells, headers):
                        cell.text = stringify(value)
                    for row in table['rows']:
                        for cell, value in zip(native.add_row().cells, row):
                            cell.text = stringify(value)
                document.save(temporary_path)
            elif kind == 'xlsx':
                from openpyxl import Workbook
                workbook = Workbook()
                workbook.remove(workbook.active)
                sheets = artifact.get('sheets', tables) or [{'name': 'Data', 'headers': ['Text'], 'rows': [[p] for p in paragraphs]}]
                for index, sheet in enumerate(sheets):
                    native = workbook.create_sheet(sheet.get('name', 'Sheet' + str(index+1)))
                    for row in [sheet.get('headers', []), *sheet.get('rows', [])]:
                        native.append(row)
                        # Untrusted strings must not become executable spreadsheet formulas.
                        for cell in native[native.max_row]:
                            if isinstance(cell.value, str) and cell.value.startswith('='):
                                cell.data_type = 's'
                from openpyxl.styles import Font, PatternFill, Alignment
                from openpyxl.utils import get_column_letter
                from math import ceil
                for native in workbook:
                    native.freeze_panes = 'A2'
                    native.auto_filter.ref = native.dimensions
                    native.print_title_rows = '1:1'
                    native.sheet_properties.pageSetUpPr.fitToPage = True
                    native.page_setup.fitToWidth = 1
                    native.page_setup.fitToHeight = 0
                    for column in native.iter_cols():
                        width = min(60, max(12, max(len(stringify(c.value)) for c in column) + 3))
                        native.column_dimensions[get_column_letter(column[0].column)].width = width
                    for row in native:
                        lines = 1
                        for cell in row:
                            cell.alignment = Alignment(vertical='top', wrap_text=True)
                            width = native.column_dimensions[cell.column_letter].width
                            lines = max(lines, sum(max(1, ceil(len(line) / max(1, width - 2))) for line in stringify(cell.value).split('\n')))
                            if cell.row == 1:
                                cell.font = Font(bold=True, color='000000')
                                cell.fill = PatternFill('solid', fgColor='E7E6E6')
                        native.row_dimensions[row[0].row].height = 16 * lines + 4
                workbook.save(temporary_path)
            elif kind == 'pdf':
                from reportlab.lib import colors
                from reportlab.lib.styles import getSampleStyleSheet
                from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle
                styles = getSampleStyleSheet()
                story = []
                if artifact.get('title'):
                    story.append(Paragraph(escape(artifact['title']), styles['Title']))
                for paragraph in paragraphs or [artifact.get('text', '')]:
                    story.extend([Paragraph(escape(str(paragraph)).replace('\n', '<br/>'), styles['BodyText']), Spacer(1, 8)])
                for table in tables:
                    rows = [table['headers'], *table['rows']]
                    native = LongTable([[Paragraph(escape(stringify(cell)), styles['BodyText']) for cell in row] for row in rows], repeatRows=1)
                    native.setStyle(TableStyle([('GRID', (0,0), (-1,-1), .3, colors.grey), ('VALIGN', (0,0), (-1,-1), 'TOP')]))
                    story.append(native)
                SimpleDocTemplate(str(temporary_path)).build(story)
            elif kind in ('tsv', 'csv'):
                if tables or artifact.get('sheets'):
                    table = (tables or artifact['sheets'])[0]
                    with temporary_path.open('w', encoding='utf-8', newline='') as stream:
                        writer = csv.writer(stream, delimiter='\t' if kind == 'tsv' else ',')
                        writer.writerow(table['headers'])
                        writer.writerows(table['rows'])
                else:
                    temporary_path.write_text(artifact.get('text', ''), encoding='utf-8')
            elif kind == 'json':
                temporary_path.write_text(json.dumps(artifact.get('data', {}), ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
            elif kind == 'md':
                temporary_path.write_text(artifact.get('text', '\n\n'.join(paragraphs)), encoding='utf-8')
            else:
                raise ValueError('Unsupported artifact kind: ' + kind)
            os.replace(temporary_path, target)
        finally:
            temporary_path.unlink(missing_ok=True)
        raw = target.read_bytes()
        rendered.append({'path': artifact['path'], 'sha256': digest(raw), 'bytes': len(raw)})
    return rendered


if __name__ == '__main__':
    try:
        request = json.load(sys.stdin)
        if request['operation'] == 'read':
            result = read_sources(request['directory'])
        elif request['operation'] == 'render':
            result = render_artifacts(request['directory'], request['artifacts'])
        else:
            raise ValueError('Unknown document operation')
        print(json.dumps(result, ensure_ascii=False))
    except Exception as error:
        print(json.dumps({'error': str(error)}))
        sys.exit(1)
