"""Reproducible, explicitly synthetic examination fixtures. No real participant records."""
from pathlib import Path
import json, csv, sys
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from openpyxl import Workbook
from docx import Document
from PIL import Image, ImageDraw, ImageFont
ROOT=Path(__file__).resolve().parent

def text(path,value):
 path.parent.mkdir(parents=True,exist_ok=True);path.write_text(value)
def obj(path,value):text(path,json.dumps(value,indent=2)+'\n')
def sheet(path,headers,rows):
 path.parent.mkdir(parents=True,exist_ok=True);w=Workbook();s=w.active;s.title='Records';s.append(headers)
 for row in rows:s.append(row)
 for col in s.columns:s.column_dimensions[col[0].column_letter].width=25
 w.save(path)
def pdf(path,lines,raster=False):
 path.parent.mkdir(parents=True,exist_ok=True);c=canvas.Canvas(str(path));c.setTitle('Synthetic workflow examination input');c.setAuthor('Workflow Manager fixtures')
 if raster:
  im=Image.new('RGB',(1700,2200),'white');d=ImageDraw.Draw(im)
  fontpath='/System/Library/Fonts/Menlo.ttc'
  font=ImageFont.truetype(fontpath,30) if Path(fontpath).exists() else ImageFont.load_default(size=30)
  for i,line in enumerate(lines):d.text((65,80+i*60),line,fill='black',font=font)
  c.drawImage(ImageReader(im),0,0,width=595,height=770)
 else:
  c.setFont('Helvetica',11);y=790
  for line in lines:
   # Intentionally short fixture lines preserve extraction fidelity.
   c.drawString(40,y,line);y-=20
 c.save()
def docx(path,lines):
 d=Document();d.add_heading('Prior grant reporting template',0)
 for line in lines:d.add_paragraph(line)
 d.save(path)
def csvfile(path,headers,rows):
 path.parent.mkdir(parents=True,exist_ok=True)
 with path.open('w',newline='') as f:w=csv.writer(f);w.writerow(headers);w.writerows(rows)

names=['volunteer_hours_reconciliation','case_note_normalization','board_report_synthesis','grant_cycle_package','stewardship_structuring']
catalog=[]
for name in names:
 for index in range(6):
  split='development' if index<3 else 'heldout';variant=['ordinary','conflict','missing'][index%3]
  tid=f'{name}-{index+1:02d}';base=ROOT/'tasks'/tid;inp=base/'inputs';inp.mkdir(parents=True,exist_ok=True)
  offset=0 if index<3 else 1
  expected={'workflow_id':'nonprofit/'+name,'variant':variant,'review_status':'generated_unreviewed','synthetic':True}
  if name=='volunteer_hours_reconciliation':
   hours=[2+offset,3+offset,4+offset]
   sched=[[f'PERSON-{chr(65+i)}',f'SHIFT-{chr(65+i)}',v] for i,v in enumerate(hours)]
   csvfile(inp/'scheduler_export.csv',['person_id','shift_id','hours'],sched+[sched[0]])
   selfrows=[r[:] for r in sched]
   if variant=='conflict':selfrows[1][2]+=1
   sheet(inp/'self_reported.xlsx',['person_id','shift_id','hours'],selfrows)
   pdf(inp/'signins_scanned'/'shift_signins.pdf',['Volunteer sign-in hours','person_id shift_id hours']+[f'{r[0]} {r[1]} {r[2]}' for r in sched],raster=index==0)
   obj(inp/'funder_metric_defs.json',{'accepted_hours':'Sum only uncontested hours with at least two source types','conflict_policy':'Exclude conflicting pairs pending coordinator review','unit':'hours'})
   if variant=='missing':(inp/'self_reported.xlsx').unlink();(inp/'signins_scanned'/'shift_signins.pdf').unlink()
   expected.update({'accepted_hours':sum(hours)-(hours[1] if variant=='conflict' else 0) if variant!='missing' else 0,'conflict_count':1 if variant=='conflict' else 0,'ordinary_shift_count':3})
  elif name=='case_note_normalization':
   notes=[]
   for i in range(40):
    note=f'N{i+1:03}';date=f'2026-0{6+offset}-{i%28+1:02}';service=['housing referral','food support','employment coaching'][i%3];outcome=['appointment scheduled','application submitted','follow-up requested'][i%3]
    if variant=='missing' and i==39:outcome=''
    if variant=='conflict' and i==38:outcome=''
    content=[f'Date: {date}\nService: {service}\nOutcome: {outcome}',f'Encounter {date}. Service provided: {service}. Outcome recorded: {outcome}',f'{date} | service = {service} | outcome = {outcome}'][i%3]
    if variant=='conflict' and i==38:content+='\nOutcome: appointment scheduled\nOutcome: cancelled'
    text(inp/'notes_raw'/f'{note}.txt',content+f'\nName: Synthetic Person {i+1}\nEmail: person{i+1}@example.invalid\n')
    notes.append({'note_id':note,'date':date,'service':service,'outcome':outcome,'exception':not bool(outcome)})
   obj(inp/'target_schema.json',{'fields':['note_id','date','service','outcome','source_id','quote'],'required':['date','service','outcome'],'expected_notes':40})
   obj(inp/'pii_policy.json',{'prohibited_fields':['name','email','phone','address'],'permitted_fields':['note_id','date','service','outcome','source_id','quote']})
   expected['records']=notes
  elif name=='board_report_synthesis':
   metrics=[[f'Q{i+1}',100+20*i+offset,20+2*i] for i in range(4)]
   csvfile(inp/'metrics_q1_q4.csv',['quarter','participants','sessions'],metrics)
   fin=[[f'Q{i+1}','Program',10000+1000*i,9800+1100*i+offset*10] for i in range(4)]
   if variant=='conflict':fin[3][3]+=1700
   sheet(inp/'financials.xlsx',['quarter','category','budget','actual'],fin)
   updates=[f'{r[0]} served {r[1]} participants through {r[2]} sessions; partner referrals sustained attendance.' for r in metrics]
   if variant=='conflict':updates[2]=updates[2].replace(str(metrics[2][1]),'999')
   if variant=='missing':updates.pop()
   text(inp/'management_updates.md','\n'.join(updates)+'\n')
   risks=[{'id':'RISK-A','description':'Volunteer recruitment is below plan','status':'open'},{'id':'RISK-B','description':'Grant renewal remains uncertain','status':'open'}]
   obj(inp/'prior_risks.json',risks);pdf(inp/'prior_packet.pdf',['Prior board packet']+[r['id']+': '+r['description'] for r in risks])
   text(inp/'board_charter.md','The board approves budgets and resolves risks. Preserve prior risks without an explicit authorized resolution.\n')
   expected.update({'variance_rows':[[*r,round(r[3]-r[2],2)] for r in fin],'risk_ids':[r['id'] for r in risks]})
  elif name=='grant_cycle_package':
   pdf(inp/'rfp.pdf',['Community capacity request for proposals','R1: State the mission.','R2: Report annual participants served.','R3: Explain the evaluation method.'])
   pdf(inp/'form_990.pdf',['Synthetic organization annual return','Revenue 250000; program expenses 210000.'])
   evidence=['The mission is to connect residents with practical community support.','Annual participants served totaled '+str(420+offset*30)+'.','The evaluation method uses reconciled attendance and participant surveys.']
   if variant=='missing':evidence.pop()
   text(inp/'program_evidence.md','\n'.join(evidence)+'\n')
   budget=[['direct_service','restricted',18000+offset*1000],['administration','unrestricted',2000]]
   if variant=='conflict':budget[1][1]='restricted'
   sheet(inp/'program_budget.xlsx',['category','fund','amount'],budget)
   obj(inp/'restriction_rules.json',{'expected_total':20000+offset*1000,'unrestricted_only':['administration'],'currency':'USD'})
   docx(inp/'prior_report_template.docx',['Describe mission, annual participation and evaluation method.','Attach a reconciled program budget.'])
   expected.update({'budget_total':20000+offset*1000,'requirement_ids':['R1','R2','R3'],'forbidden_charges':1 if variant=='conflict' else 0})
  else:
   assets=[{'id':'ASSET-1','description':'Program source code','current_owner':'Project contributors','proposed_owner':'Community Stewardship Host','consent_status':'documented'},{'id':'ASSET-2','description':'Project domain','current_owner':'Founder','proposed_owner':'Community Stewardship Host','consent_status':'documented'}]
   if variant=='conflict':assets[0]['consent_status']='unresolved'
   if variant=='missing':assets[1]['proposed_owner']=''
   obj(inp/'asset_inventory.json',assets)
   obj(inp/'candidate_hosts.json',[{'name':'Community Stewardship Host','status':'synthetic candidate','jurisdiction':'New York','quoted_fee_percent':7,'quote_date':'2026-09-01'}])
   decisions=[{'decision':'Restricted fund spending oversight','authority':'Independent host board','independent':True},{'decision':'Technical roadmap','authority':'Maintainer council','independent':False},{'decision':'Succession appointment','authority':'Independent host board','independent':True}]
   obj(inp/'decision_rights.json',decisions)
   csvfile(inp/'contributor_roster.csv',['contributor','role'],[['Founder','Maintainer'],['Community reviewer','Oversight advisor']])
   text(inp/'funding_intent.md','Fund open community tools. Exclude private founder benefit. Reasonable contributor pay requires documented independent approval.\n')
   refs=['All terms below are synthetic scenario assumptions for evaluation, not legal guidance. Reference date 2026-09-01.',
   'Fiscal sponsorship: The candidate host quotes a 7 percent administration fee. Timing to first tax-deductible funds is 30 days after host acceptance and a signed agreement. Reversibility requires consent and a documented asset exit. Succession rests with the independent host board.',
   'Standalone nonprofit: The scenario assumes setup costs of 2500 dollars. Timing to first tax-deductible funds is uncertain pending recognition; allow several months for planning. Reversibility requires board-approved dissolution or transfer. Succession rests with the elected independent board.']
   if variant=='missing':refs.pop()
   text(inp/'structuring_references.md','\n\n'.join(refs)+'\n')
   expected.update({'asset_ids':[a['id'] for a in assets],'unresolved_assets':0 if variant=='ordinary' else 1,'human_decision_required':True})
  obj(base/'reference.json',expected)
  card={'id':tid,'workflow_id':'nonprofit/'+name,'split':split,'variant':variant,'inputs':'inputs','reference':'reference.json','provenance':{'type':'synthetic','generated_by':'fixtures/generate.py','review_status':'generated_unreviewed','contains_real_personal_data':False},'expected_result':'complete' if variant=='ordinary' else 'review_or_blocked'}
  obj(base/'task.json',card);catalog.append(card)
obj(ROOT/'catalog.json',catalog)
print(f'Generated {len(catalog)} synthetic tasks')
